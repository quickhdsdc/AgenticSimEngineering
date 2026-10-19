"""报告生成 Agent —— 输出 PPTX 可行性报告(对齐公司 Simulation-Report 结构)。

结构参照样件 `*_Simulation-Report.pptx` 与 AutoForm 自带 `SimReport_PowerPoint_Template`:
  封面 → 成型性结果(判据表)→ 工艺/料片 → 结论
数据来自 QuickLink 解析出的 SimResult(料厚/减薄/FLD/分区/料片)。

注:AutoForm 的 3D 彩图(FLD/减薄云图)当前 QuickLink 未导图片,故本报告为
数据+结论草稿;需要云图时,人工在 GUI 用 AfReport 出图,或后续开启 QuickLink 图片导出。
"""
from __future__ import annotations

from datetime import date
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from ..models import FeasibilityReport

_GREEN = RGBColor(0x2E, 0x7D, 0x32)
_RED = RGBColor(0xC6, 0x28, 0x28)
_GREY = RGBColor(0x55, 0x55, 0x55)


def _txt(slide, left, top, width, height, text, size=18, bold=False, color=_GREY):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def _verdict_color(ok: bool) -> RGBColor:
    return _GREEN if ok else _RED


def build_pptx(report: FeasibilityReport, out_path: str | Path) -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    s = report.spec
    sim = report.after or report.before

    # ---------- Slide 1: 封面 ----------
    sl = prs.slides.add_slide(blank)
    _txt(sl, 0.7, 0.8, 12, 1.2, "前期评估 · 成型性仿真报告", 34, bold=True, color=RGBColor(0x1A, 0x1A, 0x1A))
    _txt(sl, 0.7, 2.2, 12, 0.5, f"Customer 客户:  {s.customer or '-'}", 18)
    _txt(sl, 0.7, 2.8, 12, 0.5, f"Part number 零件号:  {s.part_number or '-'}   {s.part_name or ''}", 18)
    _txt(sl, 0.7, 3.4, 12, 0.5, f"Material 材料:  {s.material or '-'}      Thickness 料厚:  {s.thickness_mm} mm", 18)
    _txt(sl, 0.7, 4.0, 12, 0.5, f"Date 日期:  {date.today().isoformat()}", 18)
    _txt(sl, 0.7, 6.6, 12, 0.5,
         "由 mold-agent 依据 AutoForm QuickLink 结果自动生成 · 需工程师审核", 12, color=_GREY)

    # ---------- Slide 2: 成型性判据 ----------
    sl = prs.slides.add_slide(blank)
    _txt(sl, 0.7, 0.4, 12, 0.8, "成型性 Formability", 26, bold=True, color=RGBColor(0x1A, 0x1A, 0x1A))
    if sim:
        t0 = s.thickness_mm or sim.raw.get("t0_mm", 1.2)
        thr = round(t0 * 0.70, 3)
        ok_thk = sim.min_thickness_mm >= thr
        ok_fld = sim.fld_status.value != "red"
        rows = [
            ("判据 Criterion", "值 Value", "阈值 Criterion", "状态"),
            ("最小料厚 Min thickness", f"{sim.min_thickness_mm} mm", f"≥ {thr} mm (70%)", "OK" if ok_thk else "NG"),
            ("最大减薄 Max thinning", f"{sim.thinning_ratio:.1%}", "≤ 30%", "OK" if sim.thinning_ratio <= 0.30 else "NG"),
            ("成型极限 FLD", sim.fld_status.value, "无红色 no red", "OK" if ok_fld else "NG"),
            ("开裂单元 Split elems", str(sim.raw.get("n_split_elements", "-")), "= 0", "OK" if not sim.raw.get("n_split_elements") else "NG"),
        ]
        _table(sl, rows, top=1.4)
        overall_ok = report.converged if report.after else (ok_thk and ok_fld)
        _txt(sl, 0.7, 5.2, 12, 0.7,
             f"判定 Verdict:  {'OK 可成型' if overall_ok else 'NG 存在开裂/过度减薄风险'}",
             22, bold=True, color=_verdict_color(overall_ok))
        zc = sim.raw.get("zone_counts", {})
        _txt(sl, 0.7, 6.0, 12, 0.6, f"分区统计 zones:  {zc}", 12, color=_GREY)

    # ---------- Slide 3: 工艺 / 料片 ----------
    sl = prs.slides.add_slide(blank)
    _txt(sl, 0.7, 0.4, 12, 0.8, "工艺方案 Process / 料片 Blank", 26, bold=True, color=RGBColor(0x1A, 0x1A, 0x1A))
    p = report.process
    blank_info = (sim.raw.get("blank") if sim else None) or {}
    lines = [
        f"展开料片 Developed blank:  {blank_info.get('size', '-')} mm",
        f"传送方式 Feed type:  {p.feed_type if p else '-'}",
        f"工站顺序 Stations:  {' → '.join(p.stations) if p and p.stations else '-'}",
        f"机床吨位 Press:  {p.press_tonnage if p and p.press_tonnage else '待定'}",
    ]
    for i, ln in enumerate(lines):
        _txt(sl, 0.7, 1.5 + i * 0.7, 12, 0.6, ln, 18)
    _txt(sl, 0.7, 6.2, 12, 0.6,
         "注:料片为一步成型(Crashform)展开值;排样利用率需 Strip 工序确定。", 12, color=_GREY)

    # ---------- Slide 4: 优化 / 结论 ----------
    sl = prs.slides.add_slide(blank)
    _txt(sl, 0.7, 0.4, 12, 0.8, "数模优化建议 / 结论", 26, bold=True, color=RGBColor(0x1A, 0x1A, 0x1A))
    top = 1.4
    if report.optimizations:
        _txt(sl, 0.7, top, 12, 0.5, "优化动作:", 18, bold=True)
        top += 0.6
        for i, o in enumerate(report.optimizations, 1):
            _txt(sl, 1.0, top, 11.5, 0.5, f"{i}. {o.reason}", 15)
            top += 0.5
    else:
        _txt(sl, 0.7, top, 12, 0.6,
             "本次为单次评估(未做自动优化闭环)。若 NG,建议按知识库策略放大危险区 R 角/加高翻边后重评。",
             15)
        top += 0.8
    _txt(sl, 0.7, max(top + 0.3, 5.0), 12, 1.5, f"结论 Conclusion:\n{report.conclusion or '-'}", 16, bold=True)

    out_path = Path(out_path)
    prs.save(str(out_path))
    return out_path


def _table(slide, rows, top):
    n, m = len(rows), len(rows[0])
    tbl = slide.shapes.add_table(n, m, Inches(0.7), Inches(top),
                                 Inches(12), Inches(0.4 * n)).table
    widths = [Inches(4.2), Inches(2.6), Inches(3.2), Inches(2.0)]
    for j, w in enumerate(widths[:m]):
        tbl.columns[j].width = w
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.text = str(val)
            para = cell.text_frame.paragraphs[0]
            para.font.size = Pt(13)
            para.font.bold = (i == 0)
            if j == m - 1 and i > 0:
                para.font.bold = True
                para.font.color.rgb = _GREEN if val == "OK" else _RED
